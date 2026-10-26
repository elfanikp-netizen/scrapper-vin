#!/usr/bin/env python3
"""
vin_crawler.py - For each Year / Make / Model row in an Excel sheet, search Google
with Playwright and try to pull a 17-character VIN out of the results.

Usage:
    pip install playwright openpyxl
    playwright install chromium
    python vin_crawler.py search.xlsx                 # headed browser (recommended)
    python vin_crawler.py search.xlsx --limit 10      # quick test on 10 unique vehicles
    python vin_crawler.py search.xlsx --headless      # no window (gets CAPTCHA'd sooner)

Output: search_vin.xlsx  (original columns + VIN, VIN_Source, VIN_Candidates, VIN_Status)
The script is resumable: re-run the same command and it skips rows already done.

IMPORTANT CAVEAT
A Year/Make/Model identifies millions of vehicles, not one. Whatever VIN this finds
is *a* VIN of a matching vehicle (e.g. from a dealer listing), not *the* VIN of a
specific car. Each VIN is validated (check digit + model-year character), which
removes garbage matches but cannot make the VIN "belong" to your row.
"""

import argparse
import random
import re
import sys
import time
import urllib.parse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from playwright.sync_api import Error as PWError
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

# --------------------------------------------------------------------------- #
# VIN validation
# --------------------------------------------------------------------------- #
VIN_RE = re.compile(r"(?<![A-Za-z0-9])[A-HJ-NPR-Z0-9]{17}(?![A-Za-z0-9])")  # no I, O, Q

_TRANSLIT = {
    **{str(d): d for d in range(10)},
    **dict(zip("ABCDEFGH", range(1, 9))),
    **dict(zip("JKLMN", range(1, 6))),
    "P": 7,
    "R": 9,
    **dict(zip("STUVWXYZ", range(2, 10))),
}
_WEIGHTS = [8, 7, 6, 5, 4, 3, 2, 10, 0, 9, 8, 7, 6, 5, 4, 3, 2]
_YEAR_CODES = "ABCDEFGHJKLMNPRSTVWXY123456789"  # 30-year cycle starting 1980 / 2010


def check_digit_ok(vin: str) -> bool:
    total = sum(_TRANSLIT[c] * w for c, w in zip(vin, _WEIGHTS))
    expected = "X" if total % 11 == 10 else str(total % 11)
    return vin[8] == expected


def year_matches(vin: str, year: int) -> bool:
    """10th VIN character encodes the model year (cycle of 30 years)."""
    return vin[9] == _YEAR_CODES[(year - 1980) % 30]


def extract_vins(text: str, year: int, make: str, model: str, strict_year: bool = True):
    """Return list of (vin, context_hit) sorted best-first."""
    make_l = make.lower().replace("-", " ")
    model_tok = model.lower().split()[0] if model.split() else ""
    found = {}
    for m in VIN_RE.finditer(text):
        vin = m.group(0)
        if not any(ch.isdigit() for ch in vin):
            continue
        if not check_digit_ok(vin):
            continue
        if strict_year and not year_matches(vin, year):
            continue
        ctx = text[max(0, m.start() - 300): m.end() + 300].lower()
        hit = (make_l.split()[0] in ctx) + (bool(model_tok) and model_tok in ctx)
        found[vin] = max(found.get(vin, 0), hit)
    return sorted(found.items(), key=lambda kv: -kv[1])


# --------------------------------------------------------------------------- #
# Google / Playwright helpers
# --------------------------------------------------------------------------- #
class Blocked(Exception):
    pass


def close_context(context):
    try:
        context.close()
    except PWError:
        pass


def is_blocked(page) -> bool:
    if page.is_closed():
        raise Blocked("Browser was closed while waiting for search results")
    try:
        body = page.locator("body").inner_text(timeout=2_000).casefold()
        challenge_phrases = (
            "our systems have detected unusual traffic",
            "unusual traffic from your computer network",
            "verify you are not a robot",
            "verify you are human",
            "complete the captcha",
            "solve the captcha",
        )
        if any(phrase in body for phrase in challenge_phrases):
            return True
        visible_challenge = page.locator(
            "form#captcha-form:visible, iframe[src*='recaptcha']:visible, .g-recaptcha:visible"
        )
        if visible_challenge.count():
            return True
        if "/sorry/" in page.url:
            # Google may leave the tab on /sorry/ after the interactive check
            # has cleared. A visible results container means it is safe to retry
            # the original query instead of waiting on the stale URL.
            return not page.locator("#search:visible, #rso:visible").count()
    except Exception:
        return "/sorry/" in page.url
    return False


def wait_for_human(page, headless: bool, max_wait: int = 900):
    if headless:
        raise Blocked("Google CAPTCHA in headless mode")
    print("\n  >> Google is asking for a CAPTCHA. Solve it in the browser window; "
          "the script resumes automatically...", flush=True)
    t0 = time.time()
    last_notice = t0
    while time.time() - t0 < max_wait:
        if page.is_closed():
            raise Blocked("Browser was closed during manual CAPTCHA verification")
        time.sleep(2)
        if not is_blocked(page):
            print("  >> CAPTCHA cleared, continuing.\n", flush=True)
            time.sleep(random.uniform(3, 6))
            return
        if time.time() - last_notice >= 30:
            elapsed = int(time.time() - t0)
            print(f"  >> Still waiting for manual CAPTCHA completion ({elapsed}s elapsed).", flush=True)
            last_notice = time.time()
    raise Blocked("CAPTCHA not solved in time")


def accept_consent(page):
    """Dismiss Google's cookie/consent dialog if it shows up."""
    for label in ("Accept all", "I agree", "Reject all"):
        try:
            btn = page.get_by_role("button", name=label)
            if btn.count():
                btn.first.click(timeout=2000)
                page.wait_for_load_state("domcontentloaded")
                return
        except Exception:
            pass


def google_search(page, query: str, headless: bool):
    url = "https://www.google.com/search?" + urllib.parse.urlencode(
        {"q": query, "hl": "en", "num": 10, "pws": 0}
    )
    page.goto(url, wait_until="domcontentloaded", timeout=30000)
    accept_consent(page)
    challenge_retries = 0
    while is_blocked(page):
        wait_for_human(page, headless)
        # A /sorry/ URL can contain the original /search URL in its continue
        # parameter. Always reopen the exact query after manual clearance.
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
        challenge_retries += 1
        if challenge_retries >= 2 and is_blocked(page):
            raise Blocked("Google showed another CAPTCHA after manual verification")
    try:
        page.wait_for_selector("#search, #rso, div#main", timeout=10000)
    except PWTimeout:
        pass
    text = page.inner_text("body")
    links = page.eval_on_selector_all(
        "a:has(h3)", "els => els.map(e => e.href).filter(h => h.startsWith('http'))"
    )
    seen, clean = set(), []
    for h in links:
        host = urllib.parse.urlparse(h).netloc
        if "google." in host or host in seen:
            continue
        seen.add(host)
        clean.append(h)
    return text, clean


def scan_result_page(context, url: str, timeout_ms: int = 15000) -> str:
    p = context.new_page()
    try:
        p.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
        p.wait_for_timeout(1500)
        return p.inner_text("body")
    except Exception:
        return ""
    finally:
        p.close()


def find_vin(context, page, year, make, model, follow, strict_year, headless):
    query = f'{year} {make} {model} VIN'
    text, links = google_search(page, query, headless)

    hits = extract_vins(text, year, make, model, strict_year)
    if hits:
        return hits[0][0], "Google results page", hits, "found"

    for link in links[:follow]:
        time.sleep(random.uniform(1.0, 2.5))
        body = scan_result_page(context, link)
        hits = extract_vins(body, year, make, model, strict_year)
        if hits:
            return hits[0][0], link, hits, "found"

    return "", "", [], "not_found"


# --------------------------------------------------------------------------- #
# Excel I/O
# --------------------------------------------------------------------------- #
NEW_COLS = ["VIN", "VIN_Source", "VIN_Candidates", "VIN_Status"]


def is_completed_status(status) -> bool:
    return status in {"found", "not_found"}


def open_workbook(src: Path, dst: Path):
    wb = load_workbook(dst if dst.exists() else src)
    ws = wb.active
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) + 1 for name in ("Year", "Make", "Model")}
    for name in NEW_COLS:
        if name not in header:
            ws.cell(row=1, column=ws.max_column + 1, value=name).font = Font(bold=True)
            header.append(name)
    col = {name: header.index(name) + 1 for name in header if name}
    return wb, ws, idx, col


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path, help="Excel file with Year, Make, Model columns")
    ap.add_argument("-o", "--output", type=Path, help="output xlsx (default: <input>_vin.xlsx)")
    ap.add_argument("--headless", action="store_true", help="run without a visible window")
    ap.add_argument("--follow", type=int, default=3,
                    help="if the results page has no VIN, open the top N result pages (default 3)")
    ap.add_argument("--loose", action="store_true",
                    help="accept VINs whose model-year character doesn't match the row's Year")
    ap.add_argument("--min-delay", type=float, default=6.0, help="min seconds between searches")
    ap.add_argument("--max-delay", type=float, default=14.0, help="max seconds between searches")
    ap.add_argument("--limit", type=int, default=0, help="only process the first N unique vehicles (testing)")
    ap.add_argument("--save-every", type=int, default=1, help="save workbook every N searches")
    ap.add_argument("--profile", type=Path, default=Path(".chrome_profile"),
                    help="browser profile dir (keeps cookies so you get fewer CAPTCHAs)")
    args = ap.parse_args()

    dst = args.output or args.input.with_name(args.input.stem + "_vin.xlsx")
    wb, ws, idx, col = open_workbook(args.input, dst)

    # group duplicate Year/Make/Model rows so each vehicle is searched only once
    groups, order = {}, []
    for r in range(2, ws.max_row + 1):
        y, mk, md = (ws.cell(row=r, column=idx[k]).value for k in ("Year", "Make", "Model"))
        if y is None or mk is None or md is None:
            continue
        key = (int(y), str(mk).strip(), str(md).strip())
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(r)

    def done(key):
        return all(
            is_completed_status(ws.cell(row=r, column=col["VIN_Status"]).value)
            for r in groups[key]
        )

    todo = [k for k in order if not done(k)]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(order)} unique vehicles ({ws.max_row - 1} rows); {len(todo)} left to search.")
    wb.save(dst)

    def write(key, vin, src, hits, status):
        cands = ", ".join(v for v, _ in hits[:5])
        for r in groups[key]:
            ws.cell(row=r, column=col["VIN"], value=vin)
            ws.cell(row=r, column=col["VIN_Source"], value=src)
            ws.cell(row=r, column=col["VIN_Candidates"], value=cands)
            ws.cell(row=r, column=col["VIN_Status"], value=status)

    args.profile.mkdir(exist_ok=True)
    found = 0
    with sync_playwright() as pw:
        context = pw.chromium.launch_persistent_context(
            str(args.profile),
            headless=args.headless,
            viewport={"width": 1280, "height": 900},
            locale="en-US",
        )
        page = context.pages[0] if context.pages else context.new_page()

        try:
            for n, key in enumerate(todo, 1):
                year, make, model = key
                try:
                    vin, src, hits, status = find_vin(
                        context, page, year, make, model,
                        args.follow, not args.loose, args.headless,
                    )
                except Blocked as e:
                    print(f"\nStopped: {e}. Progress saved; re-run later to resume.")
                    break
                except Exception as e:  # network hiccup, selector change, etc.
                    vin, src, hits, status = "", "", [], f"error: {type(e).__name__}"

                write(key, vin, src, hits, status)
                found += bool(vin)
                print(f"[{n}/{len(todo)}] {year} {make} {model} -> {vin or status}", flush=True)

                if n % max(args.save_every, 1) == 0:
                    wb.save(dst)
                time.sleep(random.uniform(args.min_delay, args.max_delay))
        except KeyboardInterrupt:
            print("\nInterrupted. Saving progress...")
        finally:
            wb.save(dst)
            close_context(context)

    print(f"\nSaved {dst}  ({found} VINs found this run)")


if __name__ == "__main__":
    sys.exit(main())
