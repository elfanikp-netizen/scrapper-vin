from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import Error as PlaywrightError, Page, sync_playwright


SITE_URL = "https://vininspect.com/"
CDP_PORT = 9238
VERIFICATION_TIMEOUT_SECONDS = 900
VIN_PATTERN = re.compile(r"(?<![A-HJ-NPR-Z0-9])[A-HJ-NPR-Z0-9]{17}(?![A-HJ-NPR-Z0-9])", re.IGNORECASE)
OUTPUT_HEADERS = ("Year", "Make", "Model", "VIN", "Status")


class BrowserClosedError(RuntimeError):
    pass


def slugify(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(value).strip().casefold()).strip("-")


def build_vehicle_url(year: object, make: object, model: object) -> str:
    parts = (slugify(make), slugify(model), slugify(year))
    if not all(parts):
        raise ValueError("Year, make, and model are required to build a VinInspect URL.")
    return f"{SITE_URL}vin/{parts[0]}/{parts[1]}/{parts[2]}"


def extract_vins(text: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for match in VIN_PATTERN.finditer(text.upper()):
        vin = match.group(0)
        if vin not in seen:
            found.append(vin)
            seen.add(vin)
    return found


def read_vehicles(input_path: Path, sheet_name: str | None, header_row: int) -> list[tuple[object, object, object]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet {sheet_name!r} was not found in {input_path.name}.")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        headers = {
            str(cell.value or "").strip().casefold(): index
            for index, cell in enumerate(sheet[header_row], start=1)
        }
        missing = [name for name in ("year", "make", "model") if name not in headers]
        if missing:
            raise ValueError(f"Required column(s) not found in row {header_row}: {', '.join(missing)}.")

        vehicles = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            values = tuple(row[headers[name] - 1] if headers[name] <= len(row) else None for name in ("year", "make", "model"))
            if all(value is not None and str(value).strip() for value in values):
                vehicles.append(values)
        return vehicles
    finally:
        workbook.close()


def page_needs_verification(page: Page) -> bool:
    try:
        title = page.title().casefold()
        text = page.locator("body").inner_text(timeout=3_000).casefold()
    except Exception:
        return False
    return any(
        phrase in title or phrase in text
        for phrase in (
            "just a moment",
            "security verification",
            "verify you are human",
            "verifying you are human",
            "performing security verification",
        )
    )


def wait_for_manual_verification(page: Page) -> None:
    if page.is_closed():
        raise BrowserClosedError("The browser page was closed.")
    if not page_needs_verification(page):
        return
    print(
        "VinInspect is requesting human verification. Complete it in the browser; "
        "the crawler will continue when the page is available.",
        flush=True,
    )
    deadline = time.monotonic() + VERIFICATION_TIMEOUT_SECONDS
    while True:
        if page.is_closed():
            raise BrowserClosedError("The browser page was closed during verification.")
        if not page_needs_verification(page):
            print("Verification page cleared; continuing the crawl.", flush=True)
            return
        if time.monotonic() >= deadline:
            raise TimeoutError(
                "Manual verification did not clear within 15 minutes. "
                "Complete it in the browser and rerun to resume."
            )
        try:
            page.wait_for_timeout(1_000)
        except PlaywrightError as error:
            if page.is_closed():
                raise BrowserClosedError("The browser page was closed during verification.") from error
            raise


def crawl_vehicle(page: Page, year: object, make: object, model: object) -> tuple[str, str]:
    url = build_vehicle_url(year, make, model)
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=60_000)
    except PlaywrightError as error:
        if page.is_closed():
            raise BrowserClosedError("The browser was closed during navigation.") from error
        raise
    wait_for_manual_verification(page)
    body = page.locator("body").inner_text(timeout=20_000)
    vins = extract_vins(body)
    return (vins[0], "success") if vins else ("", "vin_not_found")


def find_chrome(custom_path: Path | None = None) -> Path | None:
    if custom_path is not None:
        candidate = custom_path.expanduser().resolve()
        return candidate if candidate.is_file() else None
    candidates = [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "C:/Users/Default/AppData/Local")) / "Google/Chrome/Application/chrome.exe",
    ]
    return next((candidate for candidate in candidates if candidate.is_file()), None)


def cdp_is_ready() -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{CDP_PORT}/json/version", timeout=1):
            return True
    except Exception:
        return False


def start_chrome(chrome_path: Path, profile_path: Path) -> subprocess.Popen:
    if cdp_is_ready():
        raise RuntimeError(
            f"Chrome DevTools port {CDP_PORT} is already in use. Close the previous VinInspect crawler browser and retry."
        )
    profile_path.mkdir(parents=True, exist_ok=True)
    process = subprocess.Popen(
        [
            str(chrome_path),
            f"--remote-debugging-port={CDP_PORT}",
            f"--user-data-dir={profile_path}",
            "--no-first-run",
            "--no-default-browser-check",
            SITE_URL,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(60):
        if cdp_is_ready():
            return process
        if process.poll() is not None:
            raise RuntimeError(f"Chrome exited during startup with code {process.returncode}.")
        time.sleep(0.5)
    process.terminate()
    raise RuntimeError(f"Chrome started but DevTools did not open on port {CDP_PORT}.")


def load_previous_results(path: Path) -> dict[tuple[str, str, str], list[tuple[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [str(cell or "").strip().casefold() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        required = {name: headers.index(name) for name in ("year", "make", "model", "vin", "status") if name in headers}
        if len(required) != 5:
            return {}
        previous: dict[tuple[str, str, str], list[tuple[str, str]]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            vehicle = tuple(str(row[required[name]] or "").strip().casefold() for name in ("year", "make", "model"))
            if not all(vehicle):
                continue
            result = (str(row[required["vin"]] or "").strip(), str(row[required["status"]] or "").strip())
            previous.setdefault(vehicle, []).append(result)
        return previous
    finally:
        workbook.close()


def is_completed_result(vin: str, status: str) -> bool:
    return (status == "success" and bool(vin)) or status == "vin_not_found"


def save_results(path: Path, rows: list[tuple[object, object, object, str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(OUTPUT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 10, "B": 18, "C": 24, "D": 22, "E": 20}.items():
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f"{path.stem}.tmp{path.suffix}")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Find a VIN for each vehicle row on VinInspect.")
    parser.add_argument("--input", type=Path, default=Path(__file__).with_name("search.xlsx"), help="Input workbook (default: search.xlsx).")
    parser.add_argument("--output", type=Path, help="Output workbook (default: <input>_results.xlsx).")
    parser.add_argument("--sheet", help="Worksheet name (default: active worksheet).")
    parser.add_argument("--header-row", type=int, default=1, help="Header row number (default: 1).")
    parser.add_argument("--chrome-path", type=Path, help="Path to installed Google Chrome executable.")
    parser.add_argument("--limit", type=int, help="Process only the first N vehicle rows.")
    parser.add_argument("--start-over", action="store_true", help="Ignore existing results and crawl every row again.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    if not input_path.is_file() or input_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        print(f"Input must be an existing .xlsx or .xlsm workbook: {input_path}", file=sys.stderr)
        return 2
    if args.header_row < 1:
        print("--header-row must be 1 or greater.", file=sys.stderr)
        return 2
    if args.limit is not None and args.limit < 1:
        print("--limit must be 1 or greater.", file=sys.stderr)
        return 2
    output_path = (args.output or input_path.with_name(f"{input_path.stem}_results.xlsx")).expanduser().resolve()
    if output_path == input_path:
        print("Input and output paths must be different.", file=sys.stderr)
        return 2
    try:
        vehicles = read_vehicles(input_path, args.sheet, args.header_row)
    except (OSError, ValueError, KeyError) as error:
        print(f"Could not read input workbook: {error}", file=sys.stderr)
        return 2
    if not vehicles:
        print("No complete Year/Make/Model rows were found.", file=sys.stderr)
        return 2
    if args.limit is not None:
        vehicles = vehicles[: args.limit]

    previous = {}
    if not args.start_over and output_path.is_file():
        try:
            previous = load_previous_results(output_path)
        except (OSError, ValueError, KeyError) as error:
            print(f"Could not load existing results: {error}", file=sys.stderr)
            return 2

    results: list[tuple[object, object, object, str, str]] = []
    pending_indices = []
    for index, vehicle in enumerate(vehicles):
        key = tuple(str(value).strip().casefold() for value in vehicle)
        existing = previous.get(key, [])
        old_vin, old_status = existing.pop(0) if existing else ("", "")
        if is_completed_result(old_vin, old_status):
            results.append((*vehicle, old_vin, old_status))
        else:
            results.append((*vehicle, old_vin, "pending"))
            pending_indices.append(index)

            save_results(output_path, results)
    if not pending_indices:
        print(f"All {len(results)} vehicle rows already have VINs in {output_path}.")
        return 0
    chrome_path = find_chrome(args.chrome_path)
    if chrome_path is None:
        print("Google Chrome was not found. Install Chrome or pass --chrome-path.", file=sys.stderr)
        return 2

    print(f"Loaded {len(vehicles)} vehicle row(s); {len(pending_indices)} to crawl.")
    profile_path = Path(__file__).resolve().parent / ".browser-profile"
    chrome_process = None
    browser = None
    with sync_playwright() as playwright:
        chrome_process = start_chrome(chrome_path, profile_path)
        try:
            browser = playwright.chromium.connect_over_cdp(f"http://127.0.0.1:{CDP_PORT}")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            page = context.pages[0] if context.pages else context.new_page()
            page.set_default_timeout(30_000)
            for count, index in enumerate(pending_indices, start=1):
                year, make, model = vehicles[index]
                print(f"[{count}/{len(pending_indices)}] {year} {make} {model}", flush=True)
                try:
                    vin, status = crawl_vehicle(page, year, make, model)
                    print(f"  {vin if vin else status}", flush=True)
                except Exception as error:
                    vin = ""
                    if isinstance(error, BrowserClosedError) or page.is_closed():
                        results[index] = (year, make, model, vin, "stopped: browser closed")
                        save_results(output_path, results)
                        print("Browser closed; stopping now. Saved progress can be resumed.", flush=True)
                        break
                    status = f"error: {type(error).__name__}: {error}"
                    print(f"  {status}", flush=True)
                results[index] = (year, make, model, vin, status)
                save_results(output_path, results)
                if count < len(pending_indices):
                    time.sleep(1)
        finally:
            if browser is not None:
                try:
                    browser.close()
                except PlaywrightError:
                    pass
            if chrome_process is not None and chrome_process.poll() is None:
                chrome_process.terminate()
                try:
                    chrome_process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    chrome_process.kill()
    print(f"Saved {len(results)} row(s) to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())